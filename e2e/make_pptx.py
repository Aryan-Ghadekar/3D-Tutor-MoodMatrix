import sys
from pptx import Presentation

TEXT = [
    "Binary search trees keep keys ordered so lookup insertion and deletion take logarithmic time on balanced trees. "
    "Important: an unbalanced tree degrades to a linked list. Remember that rotations restore balance in AVL trees. ",
    "Hash tables map keys to buckets using a hash function. A good hash function spreads keys uniformly. "
    "Collisions are resolved by chaining or open addressing. Note that the load factor controls resizing. ",
    "Graph traversal uses breadth first search with a queue or depth first search with a stack. "
    "The key idea of Dijkstra algorithm is greedy selection of the closest unvisited vertex. ",
]
prs = Presentation()
for t in TEXT:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = (t * 25)
prs.save(sys.argv[1])
