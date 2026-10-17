import warnings
warnings.filterwarnings("ignore")

from fastapi import FastAPI, UploadFile, File, HTTPException, Body
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional, List
from concurrent.futures import ThreadPoolExecutor
from functools import lru_cache
import re
import uuid
import tempfile
import threading
import pdfplumber
import os
import requests
import json
from datetime import date, datetime
from dotenv import load_dotenv
from pathlib import Path
from pptx import Presentation
import random

import classroom
from rag import VectorStore

# Heavy ML dependencies (whisper, sentence-transformers, yt-dlp) are imported lazily
# so the API starts fast and the non-ML endpoints keep working if they are missing.

# =========================================================
# ENV
# =========================================================
# Load from .env file in the same directory as main.py
load_dotenv(dotenv_path=os.path.join(os.path.dirname(__file__), '.env'))

PORT = int(os.getenv("PORT", "8001"))

# =========================================================
# APP INIT
# =========================================================
app = FastAPI(title="Study Plan Generator API")

_default_origins = "http://localhost:3000,http://localhost:3001,http://127.0.0.1:3000"
app.add_middleware(
    CORSMiddleware,
    allow_origins=[o.strip() for o in os.getenv("CORS_ORIGINS", _default_origins).split(",") if o.strip()],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.include_router(classroom.router)

# =========================================================
# PYDANTIC MODELS
# =========================================================
class MCQRequest(BaseModel):
    lecture_plan: dict
    num_questions: int = 10
    question_types: List[str] = []

class AnswerAnalysisRequest(BaseModel):
    session_id: str
    questions: List[dict]
    answers: List[dict]
    student_goal: Optional[str] = None

class SaveResultsRequest(BaseModel):
    session_id: str
    lecture_id: str
    score: float
    correct: int
    total: int
    questions: List[dict]
    answers: List[dict]
    misconceptions: List[dict]
    weak_topics: List[str]
    recommendations: List[str]

class TopicQuizRequest(BaseModel):
    topic: dict
    difficulty: str = "medium"
    studentGoal: str = "gate"
    numberOfQuestions: int = 5
    sessionId: Optional[str] = None
    misconceptions: List[dict] = []
    previousPerformance: dict = {}
    cognitiveLoad: str = "medium"
    recentContext: str = ""   # e.g. the last minutes of the teacher's live audio

# =========================================================
# ML MODELS (lazy)
# =========================================================
_model_lock = threading.Lock()
_whisper_model = None
_semantic_model = None

def get_whisper_model():
    global _whisper_model
    with _model_lock:
        if _whisper_model is None:
            try:
                import whisper
            except ImportError:
                raise HTTPException(status_code=501, detail="openai-whisper is not installed; audio/YouTube transcription is unavailable.")
            _whisper_model = whisper.load_model(os.getenv("WHISPER_MODEL", "base"))
    return _whisper_model

def get_semantic_model():
    """Returns the sentence-transformers model, or None if it cannot be loaded."""
    global _semantic_model
    with _model_lock:
        if _semantic_model is None:
            try:
                from sentence_transformers import SentenceTransformer
                _semantic_model = SentenceTransformer("all-MiniLM-L6-v2")
            except Exception as e:
                print(f"Semantic model unavailable, using keyword-only scoring: {e}")
                _semantic_model = False
    return _semantic_model or None

# =========================================================
# OPENROUTER CONFIG
# =========================================================
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY")
OPENROUTER_URL = os.getenv("OPENROUTER_URL", "https://openrouter.ai/api/v1/chat/completions")
OPENROUTER_MODEL = os.getenv("OPENROUTER_MODEL", "openai/gpt-3.5-turbo")

if not OPENROUTER_API_KEY:
    print("WARNING: OPENROUTER_API_KEY not found in environment!")
    print("   Create", os.path.join(os.path.dirname(__file__), '.env'), "containing: OPENROUTER_API_KEY=your_key_here")
    print("   Without it, AI-generated content falls back to simple text extraction.")
else:
    print("OpenRouter API key loaded.")

# =========================================================
# CONSTANTS
# =========================================================
KEYWORDS = {
    "important": 1.0, "remember": 0.8, "note": 0.6, "exam": 0.9,
    "key": 0.7, "essential": 0.8, "critical": 0.9, "main": 0.7,
    "chapter": 0.85, "section": 0.7, "unit": 0.8, "concept": 0.75
}

# =========================================================
# EXTRACTION HELPERS
# =========================================================
def _save_upload(upload: UploadFile, default_suffix: str) -> str:
    suffix = Path(upload.filename or "").suffix or default_suffix
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(upload.file.read())
        return tmp.name

def extract_pdf_text(pdf_file: UploadFile) -> str:
    text = ""
    tmp_path = _save_upload(pdf_file, ".pdf")
    try:
        with pdfplumber.open(tmp_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not read the PDF file: {e}")
    finally:
        os.unlink(tmp_path)

    return text.strip()

def extract_ppt_text(ppt_file: UploadFile) -> str:
    text = ""
    tmp_path = _save_upload(ppt_file, ".pptx")
    try:
        prs = Presentation(tmp_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Could not read the PowerPoint file (only .pptx is supported): {e}")
    finally:
        os.unlink(tmp_path)

    return text.strip()

def transcribe_audio(audio_file: UploadFile) -> str:
    tmp_path = _save_upload(audio_file, ".wav")
    try:
        return get_whisper_model().transcribe(tmp_path)["text"].strip()
    finally:
        os.unlink(tmp_path)

def extract_video_id(url: str) -> Optional[str]:
    match = re.search(r"(?:v=|youtu\.be/|/embed/|/shorts/|/live/)([A-Za-z0-9_-]{11})", url)
    return match.group(1) if match else None

def process_yt_link(yt_link: str) -> str:
    """Return the transcript text of a YouTube video, or raise HTTPException."""
    video_id = extract_video_id(yt_link)
    if not video_id:
        raise HTTPException(status_code=422, detail="Could not find a video id in the YouTube link.")

    # 1) Captions (fast). Works with both the 1.x and the older youtube-transcript-api.
    try:
        from youtube_transcript_api import YouTubeTranscriptApi
        if hasattr(YouTubeTranscriptApi, "get_transcript"):
            items = YouTubeTranscriptApi.get_transcript(video_id)
            text = " ".join(item["text"] for item in items)
        else:
            fetched = YouTubeTranscriptApi().fetch(video_id)
            text = " ".join(getattr(item, "text", None) or item["text"] for item in fetched)
        if text.strip():
            return text.strip()
    except Exception as e:
        print(f"No captions for {video_id} ({type(e).__name__}); falling back to audio transcription")

    # 2) Download the audio and transcribe it with whisper (needs yt-dlp + ffmpeg).
    try:
        import yt_dlp
    except ImportError:
        raise HTTPException(status_code=501, detail="This video has no captions and yt-dlp is not installed.")

    with tempfile.TemporaryDirectory() as tmp_dir:
        ydl_opts = {
            'format': 'bestaudio/best',
            'outtmpl': os.path.join(tmp_dir, 'audio.%(ext)s'),
            'quiet': True,
            'no_warnings': True,
        }
        try:
            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                ydl.download([f"https://www.youtube.com/watch?v={video_id}"])
            files = list(Path(tmp_dir).glob("audio.*"))
            if not files:
                raise RuntimeError("download produced no audio file")
            return get_whisper_model().transcribe(str(files[0]))["text"].strip()
        except HTTPException:
            raise
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Could not transcribe the YouTube video: {e}")

# =========================================================
# TEXT PROCESSING HELPERS
# =========================================================
def chunk_text(text: str, max_tokens: int = 150):
    """Split text into chunks of roughly max_tokens words."""
    words = text.split()
    return [" ".join(words[i:i + max_tokens]) for i in range(0, len(words), max_tokens)]

def calculate_topic_count(text: str) -> int:
    """Dynamic topic count based on content length"""
    word_count = len(text.split())

    if word_count < 500:
        return 3
    elif word_count < 1500:
        return 5
    elif word_count < 3000:
        return 8
    elif word_count < 5000:
        return 12
    elif word_count < 10000:
        return 15
    else:
        return 20

def _words(text: str) -> set:
    return set(re.findall(r"[a-z0-9]{4,}", text.lower()))

def extract_features(chunks, transcript):
    model = get_semantic_model()
    embeddings = model.encode(chunks) if model is not None else None
    doc_embedding = embeddings.mean(axis=0) if embeddings is not None else None

    # Words that appear in "important" sentences of the spoken transcript
    spoken_key_words = set()
    if transcript:
        for sentence in re.split(r"[.!?]", transcript):
            if any(k in sentence.lower() for k in KEYWORDS):
                spoken_key_words |= _words(sentence)

    features = []
    for i, chunk in enumerate(chunks):
        if embeddings is not None:
            from sentence_transformers import util
            semantic = float(util.cos_sim(embeddings[i], doc_embedding)[0][0])
        else:
            semantic = 0.5
        keyword = min(sum(v for k, v in KEYWORDS.items() if k in chunk.lower()), 1.0)
        duration = min(len(chunk.split()) / 150, 1.0)

        audio = 0.0
        if spoken_key_words:
            chunk_words = _words(chunk)
            if chunk_words:
                audio = min(len(chunk_words & spoken_key_words) / max(len(chunk_words) * 0.3, 1), 1.0)

        features.append({
            "semantic": semantic,
            "keyword": keyword,
            "duration": duration,
            "audio": audio
        })

    return features

def compute_importance(features):
    return [
        0.35 * f["semantic"]
        + 0.25 * f["audio"]
        + 0.20 * f["duration"]
        + 0.20 * f["keyword"]
        for f in features
    ]

def label_importance(score):
    if score >= 0.70:
        return "high"
    elif score >= 0.45:
        return "medium"
    return "low"

# =========================================================
# API VALIDATION HELPERS
# =========================================================
def validate_api_key():
    """Validate that API key is set and non-empty"""
    if not OPENROUTER_API_KEY or OPENROUTER_API_KEY.strip() == "":
        raise Exception(
            "OPENROUTER_API_KEY is not configured. "
            "Please set it in backend/fastapi-service/.env file. "
            "Format: OPENROUTER_API_KEY=sk-or-v1-xxxxx"
        )
    return True

def make_openrouter_request(payload, timeout=30):
    """Make OpenRouter API request with proper error handling"""
    validate_api_key()

    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:3000",
        "X-Title": "Study Session AI"
    }

    try:
        response = requests.post(OPENROUTER_URL, json=payload, headers=headers, timeout=timeout)

        if response.status_code == 401:
            raise Exception("OpenRouter API returned 401 Unauthorized. Please verify your API key is correct and active.")

        response.raise_for_status()
        return response.json()

    except requests.exceptions.Timeout:
        raise Exception(f"OpenRouter API request timed out after {timeout}s. Try again.")
    except requests.exceptions.ConnectionError as e:
        raise Exception(f"Cannot connect to OpenRouter API: {str(e)}")

def llm_message(payload, timeout=30) -> str:
    """Send a chat payload and return the reply text."""
    result = make_openrouter_request(payload, timeout=timeout)
    choices = result.get("choices") or []
    if not choices or not choices[0].get("message"):
        raise Exception("Invalid API response format")
    return choices[0]["message"]["content"].strip()

def parse_llm_json(content: str):
    """Parse JSON from an LLM reply, tolerating code fences and surrounding prose."""
    content = content.strip().replace("```json", "").replace("```", "").strip()
    try:
        return json.loads(content)
    except json.JSONDecodeError:
        pass
    for open_ch, close_ch in (("{", "}"), ("[", "]")):
        start, end = content.find(open_ch), content.rfind(close_ch)
        if start != -1 and end > start:
            try:
                return json.loads(content[start:end + 1])
            except json.JSONDecodeError:
                continue
    raise ValueError("LLM reply did not contain valid JSON")

# =========================================================
# LLM GENERATION HELPERS
# =========================================================
def first_sentences(text: str, n: int = 2, limit: int = 300) -> str:
    sentences = re.split(r"(?<=[.!?])\s+", text.strip())
    return " ".join(sentences[:n])[:limit]

def generate_topic_title(text):
    """Generate concise, meaningful topic titles"""
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {"role": "system", "content": "You are an expert educator. Generate a clear, concise topic title (max 10 words) that captures the main concept. Focus on key technical terms and concepts. Reply with the title only."},
            {"role": "user", "content": f"Generate a topic title for this content:\n\n{text[:1000]}"}
        ],
        "temperature": 0.3,
        "max_tokens": 40
    }

    try:
        title = llm_message(payload).replace('"', '').strip().rstrip('.')
        if title:
            return title[:100]
    except Exception as e:
        print(f"Title generation failed: {e}")
    return " ".join(text.split()[:8])

def generate_summary(text):
    """Generate concise summary instead of full chunk"""
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {"role": "system", "content": "Summarize the following content in 2-3 clear, informative sentences. Focus on key concepts and technical details."},
            {"role": "user", "content": text[:1500]}
        ],
        "temperature": 0.4,
        "max_tokens": 150
    }

    try:
        summary = llm_message(payload)
        if summary:
            return summary
    except Exception as e:
        print(f"Summary generation failed: {e}")
    return first_sentences(text, 3, 400)

def generate_structured_content(prompt, text):
    """Ask the LLM for a short JSON list of strings. Returns [] on any failure."""
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {"role": "system", "content": f"{prompt}\n\nIMPORTANT: Return ONLY a JSON array of strings, nothing else. Example: [\"item 1\", \"item 2\", \"item 3\"]"},
            {"role": "user", "content": text[:1200]}
        ],
        "temperature": 0.5,
        "max_tokens": 200
    }

    try:
        content = llm_message(payload)
        try:
            parsed = parse_llm_json(content)
            if isinstance(parsed, list):
                return [str(p) for p in parsed if str(p).strip()][:4]
            return [str(parsed)]
        except ValueError:
            items = [line.strip("-• ").strip() for line in content.split("\n") if line.strip()]
            return [item for item in items if item][:4]
    except Exception as e:
        print(f"Structured content generation failed: {e}")
        return []

# =========================================================
# ANALYSIS FUNCTIONS
# =========================================================
def analyze_text(source_text, transcript):
    chunks = chunk_text(source_text)
    if not chunks:
        raise HTTPException(status_code=422, detail="No readable text was found in the provided files or video.")

    features = extract_features(chunks, transcript)
    scores = compute_importance(features)

    analyzed = [
        {"index": i, "text": c, "score": scores[i], "importance": label_importance(scores[i])}
        for i, c in enumerate(chunks)
    ]

    # Keep the highest-scoring chunks, but present them in lecture order
    topic_count = calculate_topic_count(source_text)
    top = sorted(analyzed, key=lambda x: x["score"], reverse=True)[:topic_count]
    top.sort(key=lambda x: x["index"])

    return {
        "chunks": top,
        "all_chunks": chunks,
        "has_audio": transcript is not None,
        "total_chunks": len(analyzed),
        "word_count": len(source_text.split())
    }

def build_topic(i, c):
    text = c["text"]
    with ThreadPoolExecutor(max_workers=5) as pool:
        f_title = pool.submit(generate_topic_title, text)
        f_summary = pool.submit(generate_summary, text)
        f_prereq = pool.submit(generate_structured_content, "List 2-3 key prerequisites or background knowledge needed to understand this topic", text)
        f_misc = pool.submit(generate_structured_content, "List 2-3 common misconceptions or mistakes students make about this topic", text)
        f_insights = pool.submit(generate_structured_content, "List 2-3 key insights, takeaways, or important points from this topic", text)

        return {
            "id": f"topic-{i+1}",
            "title": f_title.result(),
            "description": f_summary.result(),
            "full_content": text[:500],
            "importance": c["importance"],
            "weightage": round(c["score"] * 100, 2),
            "prerequisites": f_prereq.result(),
            "commonMisconceptions": f_misc.result(),
            "keyInsights": f_insights.result()
        }

def build_lecture_plan(data):
    """Build comprehensive lecture plan with dynamic topics"""
    chunks = data["chunks"][:20]
    print(f"Generating {len(chunks)} topics from {data.get('word_count', 0)} words...")

    topics = [build_topic(i, c) for i, c in enumerate(chunks)]

    all_text = " ".join(c["text"] for c in chunks[:5])

    overall_topic = generate_topic_title(all_text) if all_text else "Lecture"
    duration = (generate_structured_content(
        "Estimate lecture duration based on content complexity. Return a JSON array with a single string like \"45-60 minutes\"",
        all_text) or ["30-45 minutes"])[0]
    instructor = (generate_structured_content(
        "Suggest appropriate instructor role/expertise for this content. Return a JSON array with a single string like \"Computer Architecture Professor\"",
        all_text) or ["Subject Matter Expert"])[0]

    return {
        "sessionId": str(uuid.uuid4()),
        "lectureId": str(uuid.uuid4()),
        "topic": overall_topic,
        "duration": duration,
        "date": str(date.today()),
        "instructor": instructor,
        "impTopics": topics,
        "metadata": {
            "analysisDepth": "signal-weighted-enhanced",
            "confidenceScore": round(
                sum(c["score"] for c in chunks) / max(len(chunks), 1), 2
            ),
            "fallbackUsed": not data["has_audio"],
            "totalTopicsGenerated": len(topics),
            "wordCount": data.get("word_count", 0),
            "processingVersion": "2.1"
        }
    }

# =========================================================
# MCQ GENERATION
# =========================================================
OPTION_IDS = ["A", "B", "C", "D"]

def normalize_mcq(mcq):
    """Validate an LLM MCQ and shuffle its options so the answer position is random.

    Returns a dict with question/options/explanation, or None if the MCQ is unusable.
    """
    if not isinstance(mcq, dict) or not str(mcq.get("question", "")).strip():
        return None
    raw = mcq.get("options")
    if not isinstance(raw, list) or len(raw) != 4:
        return None

    options = []
    for opt in raw:
        if not isinstance(opt, dict) or not str(opt.get("text", "")).strip():
            return None
        options.append({"text": str(opt["text"]).strip(), "isCorrect": bool(opt.get("isCorrect"))})

    if sum(o["isCorrect"] for o in options) != 1:
        return None

    random.shuffle(options)
    for i, opt in enumerate(options):
        opt["id"] = OPTION_IDS[i]
    options = [{"id": o["id"], "text": o["text"], "isCorrect": o["isCorrect"]} for o in options]

    return {
        "question": str(mcq["question"]).strip(),
        "options": options,
        "explanation": str(mcq.get("explanation", "")).strip(),
    }

MCQ_SYSTEM_PROMPT = """You are an expert educator creating multiple choice questions.
Generate questions with 4 options where:
- The question tests understanding of the concept, not trivia
- Options are plausible and similar in length
- Only ONE option is correct
- Wrong options are believable but incorrect

Return ONLY valid JSON in this exact format:
{
  "question": "Question text here?",
  "options": [
    {"text": "First option", "isCorrect": true},
    {"text": "Second option", "isCorrect": false},
    {"text": "Third option", "isCorrect": false},
    {"text": "Fourth option", "isCorrect": false}
  ],
  "explanation": "Why the correct answer is correct"
}"""

def generate_mcq_with_llm(topic_title, content, insights):
    """Generate a single MCQ with the LLM. Returns a normalized MCQ or None."""
    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {"role": "system", "content": MCQ_SYSTEM_PROMPT},
            {"role": "user", "content": f"""Topic: {topic_title}

Content: {content[:800]}

Key Points: {', '.join(insights[:2]) if insights else 'See content above'}

Generate a multiple choice question testing this concept."""}
        ],
        "temperature": 0.7,
        "max_tokens": 400
    }

    try:
        return normalize_mcq(parse_llm_json(llm_message(payload, timeout=30)))
    except Exception as e:
        print(f"LLM MCQ generation failed: {e}")
        return None

def _candidate_statements(topic):
    """Short factual statements taken from a topic, used for template questions."""
    statements = [s.strip() for s in topic.get("keyInsights", []) if isinstance(s, str) and len(s.strip()) > 10]
    description = first_sentences(topic.get("description") or topic.get("full_content") or "", 1, 200)
    if len(description) > 10:
        statements.append(description)
    return statements

def create_diagnostic_question(qid, topic, content, insights, other_topics=()):
    """Create a diagnostic MCQ: LLM-generated when possible, otherwise template-based."""
    title = topic.get('title', 'Topic')

    llm_mcq = generate_mcq_with_llm(title, content, insights)
    if llm_mcq:
        correct_answer = next(o["id"] for o in llm_mcq["options"] if o["isCorrect"])
        return {
            "id": f"q{qid}",
            "question": llm_mcq["question"],
            "options": llm_mcq["options"],
            "correctAnswer": correct_answer,
            "explanation": llm_mcq["explanation"] or f"Based on the lecture content about {title}",
            "difficulty": topic.get('importance', 'medium'),
            "conceptsTested": [title],
            "topicId": topic.get('id'),
            "topicName": title,
            "diagnosticFocus": "knowledge_verification",
            "source": "ai"
        }

    # Template fallback: the right answer comes from this topic, distractors from other topics
    own = _candidate_statements(topic)
    correct = own[0] if own else f"{title} is covered in this lecture."
    distractors = []
    for other in other_topics:
        if other.get("id") == topic.get("id"):
            continue
        for statement in _candidate_statements(other):
            if statement != correct and statement not in distractors:
                distractors.append(statement)
                break
        if len(distractors) == 3:
            break
    generic = [
        f"{title} is not part of this lecture's material.",
        f"The opposite of this statement is true for {title}.",
        f"{title} only applies in situations the lecture did not cover.",
    ]
    for g in generic:
        if len(distractors) == 3:
            break
        distractors.append(g)

    options = [{"text": correct, "isCorrect": True}] + [{"text": d, "isCorrect": False} for d in distractors[:3]]
    random.shuffle(options)
    for i, opt in enumerate(options):
        opt["id"] = OPTION_IDS[i]
    options = [{"id": o["id"], "text": o["text"], "isCorrect": o["isCorrect"]} for o in options]

    return {
        "id": f"q{qid}",
        "question": f"Which statement is correct about {title}?",
        "options": options,
        "correctAnswer": next(o["id"] for o in options if o["isCorrect"]),
        "explanation": f"According to the lecture: {correct}",
        "difficulty": topic.get('importance', 'medium'),
        "conceptsTested": [title],
        "topicId": topic.get('id'),
        "topicName": title,
        "diagnosticFocus": "knowledge_verification",
        "source": "template"
    }

def generate_topic_quiz_with_llm(req: "TopicQuizRequest"):
    """Generate a multi-question quiz for one topic. Returns a quiz dict or None."""
    topic = req.topic
    title = topic.get("title", "Topic")
    count = max(1, min(req.numberOfQuestions, 10))
    weak = [m.get("concept") or m.get("topic") for m in req.misconceptions[:3] if isinstance(m, dict)]

    system = f"""You are an expert {req.studentGoal.upper()} exam tutor writing a multiple choice quiz.
Write exactly {count} questions at {req.difficulty} difficulty about the topic. The student's current cognitive load is {req.cognitiveLoad}; keep wording clear.
Vary the type of question: conceptual, application and analysis.
{('Target these known weak areas where relevant: ' + ', '.join(w for w in weak if w)) if any(weak) else ''}

Return ONLY valid JSON:
{{"questions": [{{"question": "...", "type": "conceptual", "options": [{{"text": "...", "isCorrect": true}}, {{"text": "...", "isCorrect": false}}, {{"text": "...", "isCorrect": false}}, {{"text": "...", "isCorrect": false}}], "explanation": "...", "misconception": "the wrong belief a wrong answer reveals"}}]}}
Each question has exactly 4 options and exactly one correct option."""
    user = f"""Topic: {title}
Description: {topic.get('description', '')[:600]}
Key insights: {'; '.join(topic.get('keyInsights', [])[:4])}
Common misconceptions: {'; '.join(topic.get('commonMisconceptions', [])[:3])}"""
    if req.recentContext.strip():
        user += f"\nThe teacher just said (base some questions on this): {req.recentContext.strip()[:1200]}"

    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}],
        "temperature": 0.6,
        "max_tokens": 2200
    }

    try:
        parsed = parse_llm_json(llm_message(payload, timeout=45))
        raw_questions = parsed.get("questions") if isinstance(parsed, dict) else parsed
    except Exception as e:
        print(f"LLM quiz generation failed: {e}")
        return None

    questions = []
    for raw in raw_questions or []:
        mcq = normalize_mcq(raw)
        if not mcq:
            continue
        n = len(questions) + 1
        qtype = raw.get("type") if raw.get("type") in ("conceptual", "application", "analysis") else "conceptual"
        questions.append({
            "id": f"q{n}_{uuid.uuid4().hex[:6]}",
            "text": mcq["question"],
            "question": mcq["question"],
            "type": qtype,
            "difficulty": req.difficulty,
            "options": mcq["options"],
            "correctAnswer": next(o["id"] for o in mcq["options"] if o["isCorrect"]),
            "explanation": mcq["explanation"],
            "misconception": str(raw.get("misconception", "")).strip(),
            "conceptsTested": [title],
            "timeEstimate": 75,
            "topic": {"id": topic.get("id"), "title": title},
        })

    if not questions:
        return None

    return {
        "id": f"ai_quiz_{uuid.uuid4().hex[:10]}",
        "title": f"{req.studentGoal.upper()} Quiz: {title}",
        "topic": title,
        "studentGoal": req.studentGoal,
        "difficulty": req.difficulty,
        "questions": questions,
        "totalQuestions": len(questions),
        "estimatedTime": sum(q["timeEstimate"] for q in questions),
        "metadata": {
            "generatedAt": datetime.now().isoformat(),
            "source": "ai",
            "topicId": topic.get("id"),
            "sessionId": req.sessionId,
        }
    }

# =========================================================
# STORAGE (JSON files under ./data, git-ignored)
# =========================================================
DATA_DIR = Path(os.getenv("DATA_DIR", Path(__file__).resolve().parent / "data"))
LECTURE_PLANS_FILE = DATA_DIR / "lecture_plans.json"
QUIZ_RESULTS_FILE = DATA_DIR / "quiz_results.json"
_file_lock = threading.Lock()

# Retrieval index: lecture chunks, transcripts and previously generated explanations
rag_store = VectorStore(DATA_DIR / "rag", model_getter=lambda: get_semantic_model())


def snippet(text: str, limit: int = 220) -> str:
    text = " ".join(text.split())
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + "..."


def index_lecture(plan: dict, analysis: dict, transcript: Optional[str], transcript_is_source: bool = False):
    """Index every chunk of the lecture (not just the ones that became topics) for retrieval.

    The transcript is indexed separately as kind "transcript" unless it already *is* the lecture text
    (transcript-only uploads), in which case indexing it again would only create duplicate hits.
    """
    topic_by_chunk = {c["index"]: f"topic-{i + 1}" for i, c in enumerate(analysis["chunks"][:20])}
    docs = [
        {"id": f"chunk-{i}", "text": text, "kind": "lecture",
         "meta": {"chunkIndex": i, "topicId": topic_by_chunk.get(i)}}
        for i, text in enumerate(analysis["all_chunks"])
    ]
    if transcript and not transcript_is_source:
        docs += [{"id": f"transcript-{i}", "text": text, "kind": "transcript", "meta": {"chunkIndex": i}}
                 for i, text in enumerate(chunk_text(transcript))]
    rag_store.add(plan["sessionId"], docs)


def retrieve_reference(collection: Optional[str], query: str, include_explanations: bool):
    """Fetch source passages (and earlier explanations) to ground a generated answer."""
    if not collection:
        return [], []
    sources = rag_store.query(collection, query, k=4, kinds=["lecture", "transcript"])
    earlier = rag_store.query(collection, query, k=2, kinds=["explanation"]) if include_explanations else []
    return sources, earlier

def _read_json_list(path: Path) -> list:
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8") or "[]")
    except (json.JSONDecodeError, OSError) as e:
        print(f"Could not read {path.name} ({e}); starting fresh")
        return []
    return data if isinstance(data, list) else [data]

def _prepend_json_list(path: Path, item: dict, keep: int) -> None:
    """Prepend an item to a JSON list file, atomically."""
    with _file_lock:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        items = [item] + _read_json_list(path)
        tmp = path.with_suffix(".tmp")
        tmp.write_text(json.dumps(items[:keep], indent=2, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, path)

def save_lecture_plan(lecture_plan):
    lecture_plan["generated_at"] = datetime.now().isoformat()
    _prepend_json_list(LECTURE_PLANS_FILE, lecture_plan, keep=10)
    print(f"Saved lecture plan with {lecture_plan['metadata']['totalTopicsGenerated']} topics")

def save_quiz_results(results_data):
    results_data["saved_at"] = datetime.now().isoformat()
    _prepend_json_list(QUIZ_RESULTS_FILE, results_data, keep=50)
    print("Saved quiz results")

# =========================================================
# AI-POWERED SLIDE GENERATION WITH OPENROUTER
# =========================================================

def generate_slide_with_ai(
    topic: dict,
    action_type: str,
    student_goal: str,
    knowledge_state: dict,
    user_question: str = None,
    context: str = "",
    reference: str = ""
) -> dict:
    """
    Generate complete slide content using AI based on:
    1. Student Goal & Exam
    2. Knowledge State from Quiz Results  
    3. Lecture Plan Content
    """
    
    knowledge_level = knowledge_state.get('knowledgeLevel', 'intermediate')
    avg_score = knowledge_state.get('averageScore', 0)
    weak_topics = knowledge_state.get('weakTopics', [])
    misconceptions = [m for m in knowledge_state.get('misconceptions', []) if isinstance(m, dict)]
    weak_topics = [str(t) for t in weak_topics]
    action_type = action_type or 'concept_explanation'
    
    # Build comprehensive AI prompt
    system_prompt = f"""You are an expert educational content creator specializing in {student_goal.upper()} exam preparation.

Your task is to create engaging, pedagogically sound learning slides that adapt to the student's knowledge level and learning needs.

CRITICAL REQUIREMENTS:
1. Generate content specifically tailored for {student_goal.upper()} exam preparation
2. Adapt complexity based on student's knowledge level: {knowledge_level} (average score: {avg_score}%)
3. Address known weak areas: {', '.join(weak_topics[:3]) if weak_topics else 'None identified yet'}
4. Correct known misconceptions: {len(misconceptions)} identified
5. Include exam-relevant examples, analogies, and real-world applications
6. Provide clear explanations with proper depth for the knowledge level
7. Add memory hooks and learning strategies

RESPONSE FORMAT (JSON):
{{
  "title": "Clear, concise slide title (max 10 words)",
  "narrative": "Engaging 3-4 paragraph explanation with examples and context. Make it conversational and educational.",
  "keyInsights": [
    "4-5 critical points the student must understand",
    "Each insight should be actionable and exam-relevant",
    "Focus on concepts frequently tested in {student_goal.upper()}"
  ],
  "analogies": [
    {{
      "domain": "Relatable real-world domain",
      "comparison": "How this concept is similar to the domain",
      "lesson": "What the student should learn from this analogy"
    }}
  ],
  "realWorldScenario": {{
    "situation": "Practical scenario where this concept applies",
    "application": "How to apply the concept in this scenario",
    "outcome": "Result of proper application"
  }},
  "interactiveElement": "Thought-provoking question to engage the student",
  "examRelevance": "Specific details about how this appears in {student_goal.upper()} exams - question types, weightage, frequency",
  "commonTraps": [
    "3-4 common mistakes students make with this concept",
    "Include misconceptions and wrong approaches"
  ],
  "memoryHook": "Memorable phrase, acronym, or pattern to remember this concept"
}}

ADAPTATION RULES:
- If knowledge_level is 'beginner': Use simple language, more examples, step-by-step breakdowns
- If knowledge_level is 'intermediate': Balance theory with practice, moderate complexity
- If knowledge_level is 'advanced': Include technical depth, edge cases, advanced applications
- If weak_topics mentioned: Provide extra clarity on related concepts
- If misconceptions exist: Explicitly address and correct them

Return ONLY valid JSON, no markdown formatting, no extra text."""

    # Build user prompt based on action type
    if action_type == 'concept_explanation':
        user_prompt = f"""Create an educational slide explaining this concept:

TOPIC: {topic.get('title', 'Unknown Topic')}

TOPIC DETAILS:
Description: {topic.get('description', topic.get('full_content', 'No description available'))[:500]}

Key Insights from Material: {', '.join(topic.get('keyInsights', [])[:3])}
Common Misconceptions: {', '.join(topic.get('commonMisconceptions', [])[:2])}
Real-World Applications: {', '.join(topic.get('realWorldApplications', [])[:2])}
Importance Level: {topic.get('importance', 'medium')}
Difficulty: {topic.get('difficulty', 'medium')}

STUDENT CONTEXT:
- Preparing for: {student_goal.upper()}
- Current knowledge level: {knowledge_level} ({avg_score}% average score)
- Weak areas: {', '.join(weak_topics[:3]) if weak_topics else 'None yet'}
- Recent quiz performance: {"Below expectations" if avg_score < 60 else "Good" if avg_score < 80 else "Excellent"}

Create a comprehensive, engaging slide that helps the student master this concept for their {student_goal.upper()} exam."""

    elif action_type == 'explain-more':
        user_prompt = f"""The student asked for MORE DETAILED explanation of:

TOPIC: {topic.get('title', 'Unknown Topic')}

CONTEXT: They've seen the basic explanation but want deeper insights.

STUDENT PROFILE:
- Knowledge level: {knowledge_level} ({avg_score}% average)
- They're ready for: {"More advanced content" if avg_score >= 70 else "Clearer explanations with more examples"}

Create a DEEPER, MORE DETAILED slide that:
1. Goes beyond surface-level understanding
2. Includes advanced insights and nuances
3. Provides multiple perspectives and approaches
4. Covers edge cases and special scenarios
5. Includes complex examples and applications
6. Addresses subtle points often missed by students

Make this significantly more comprehensive than a basic explanation."""

    elif action_type == 'ask-question':
        user_prompt = f"""The student asked this SPECIFIC QUESTION:

QUESTION: "{user_question}"

ABOUT TOPIC: {topic.get('title', 'Unknown Topic')}

STUDENT CONTEXT:
- Knowledge level: {knowledge_level}
- Current score: {avg_score}%
- Preparing for: {student_goal.upper()}

Create a slide that:
1. DIRECTLY answers their specific question
2. Provides step-by-step explanation
3. Includes concrete examples
4. Shows how to verify the answer
5. Connects to exam patterns in {student_goal.upper()}
6. Addresses related concepts they should know

Focus on making the answer clear, complete, and exam-relevant."""

    elif action_type == 'quiz-results':
        user_prompt = f"""Create a QUIZ PERFORMANCE ANALYSIS slide:

TOPIC TESTED: {topic.get('title', 'Quiz Topic')}

QUIZ CONTEXT: {context}

STUDENT PERFORMANCE:
- Current knowledge level: {knowledge_level}
- Average score: {avg_score}%
- Weak topics identified: {', '.join(weak_topics[:5])}
- Misconceptions found: {len(misconceptions)}

Recent Mistakes:
{chr(10).join([f"- {m.get('concept') or m.get('topic', 'Unknown')}: {m.get('type', 'Error')}" for m in misconceptions[:3]])}

Create a slide that:
1. Analyzes their performance objectively
2. Identifies specific areas needing improvement
3. Explains WHY they made mistakes (conceptual gaps, misconceptions, etc.)
4. Provides ACTIONABLE recommendations
5. Offers encouragement and learning strategies
6. Maps improvement areas to {student_goal.upper()} exam requirements

Be honest but supportive. Focus on growth opportunities."""

    elif action_type == 'quiz-feedback':
        # Parse context if it's JSON
        context_data = {}
        if context:
            try:
                context_data = json.loads(context) if isinstance(context, str) else context
            except:
                context_data = {}
        
        wrong_answer = context_data.get('wrongAnswer', {})
        student_performance = context_data.get('studentPerformance', {})
        
        user_prompt = f"""Create a WRONG ANSWER FEEDBACK slide for this quiz mistake:

QUESTION: "{wrong_answer.get('question', 'Unknown question')}"

STUDENT'S ANSWER: "{wrong_answer.get('selectedAnswer', {}).get('text', 'Unknown answer')}"

CORRECT ANSWER: "{wrong_answer.get('correctAnswer', {}).get('text', 'Unknown correct answer')}"

EXPLANATION: "{wrong_answer.get('explanation', 'No explanation available')}"

STUDENT PERFORMANCE:
- Current score: {avg_score}%
- Total wrong answers: {student_performance.get('wrongAnswersCount', 0)}

STUDENT CONTEXT:
- Knowledge level: {knowledge_level}
- Preparing for: {student_goal.upper()}

Create a slide that:
1. Clearly explains WHY the student's answer was wrong
2. Provides the CORRECT answer with detailed explanation
3. Addresses the underlying misconception or knowledge gap
4. Connects to the original slide content they studied
5. Offers strategies to avoid similar mistakes
6. Includes exam-relevant tips for {student_goal.upper()}

Be supportive but clear. Focus on learning from the mistake."""

    elif action_type == 'simplified-explanation':
        user_prompt = f"""The student is overloaded and struggling with this topic. Re-explain it so a tutor could say it aloud.

TOPIC: {topic.get('title', 'Unknown Topic')}
TOPIC DETAILS: {topic.get('description', topic.get('full_content', ''))[:500]}
WHAT THE TEACHER SAID IN THE LAST FEW MINUTES: {context or 'not available'}

STUDENT CONTEXT: {knowledge_level} level, preparing for {student_goal.upper()}

Rules:
1. Use short sentences and everyday words; one core idea in the narrative
2. Give one concrete real-world example and one simple analogy
3. End with a different phrasing of the key idea ("In other words...")
4. Keep keyInsights to 3 very short points"""

    elif action_type == 'completion':
        user_prompt = f"""The student has finished the whole lecture. Create a COURSE COMPLETION SUMMARY slide.

LECTURE TOPICS COVERED: {context or topic.get('title', 'the lecture')}

STUDENT CONTEXT:
- Knowledge level: {knowledge_level} ({avg_score}% average score)
- Weak areas: {', '.join(weak_topics[:3]) if weak_topics else 'None identified'}
- Preparing for: {student_goal.upper()}

Create a slide that congratulates the student, recaps the most important ideas across the topics, lists what to revise next, and gives an exam-focused study plan for {student_goal.upper()}."""

    else:
        user_prompt = f"""Create an educational slide about: {topic.get('title', 'Topic')}

For {student_goal.upper()} exam preparation at {knowledge_level} level."""

    if reference:
        user_prompt += (
            "\n\nSOURCE MATERIAL FROM THE TEACHER'S LECTURE (ground your answer in this; "
            "if it does not cover the point, say so instead of guessing):\n" + reference
        )

    payload = {
        "model": OPENROUTER_MODEL,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ],
        "temperature": 0.7,
        "max_tokens": 2000
    }

    print(f"Calling AI for {action_type}...")
    try:
        slide_content = parse_llm_json(llm_message(payload, timeout=45))
        if not isinstance(slide_content, dict):
            raise ValueError("expected a JSON object")
    except requests.exceptions.HTTPError as e:
        raise HTTPException(status_code=502, detail=f"AI model unavailable: {e}")
    except Exception as e:
        raise HTTPException(status_code=502, detail=f"Slide generation error: {e}")

    print(f"AI generated: {slide_content.get('title', 'Untitled')}")
    return slide_content


# =========================================================
# API ENDPOINTS
# =========================================================
# Note: endpoints that do blocking work (LLM calls, file/ML processing) are plain `def`
# so FastAPI runs them in its thread pool instead of blocking the event loop.

COMMONS_API = "https://commons.wikimedia.org/w/api.php"
COMMONS_HEADERS = {"User-Agent": "MoodMatrixTutor/1.0 (educational project; contact: project maintainers)"}
IMAGE_MIMES = {"image/jpeg", "image/png", "image/gif", "image/svg+xml"}

@lru_cache(maxsize=256)
def search_commons_image(query: str, stems: tuple = ()):
    """Find one relevant image on Wikimedia Commons. Returns a dict or None.

    `stems` are word prefixes of the topic; a result is only accepted if its file name contains one.
    """
    params = {
        "action": "query", "format": "json", "generator": "search",
        "gsrsearch": f"{query} filetype:bitmap|drawing", "gsrnamespace": 6, "gsrlimit": 6,
        "prop": "imageinfo", "iiprop": "url|mime|size", "iiurlwidth": 800,
    }
    try:
        r = requests.get(COMMONS_API, params=params, headers=COMMONS_HEADERS, timeout=8)
        r.raise_for_status()
        pages = sorted((r.json().get("query", {}).get("pages", {})).values(), key=lambda p: p.get("index", 0))
    except Exception as e:
        print(f"Wikimedia search failed for '{query}': {e}")
        return None

    for page in pages:
        info = (page.get("imageinfo") or [{}])[0]
        file_name = page.get("title", "").lower().replace("_", " ")
        relevant = not stems or any(stem in file_name for stem in stems)
        if relevant and info.get("mime") in IMAGE_MIMES and info.get("width", 0) >= 300 and info.get("thumburl"):
            return {"url": info["thumburl"], "title": page.get("title", ""), "page": info.get("descriptionurl", "")}
    return None

def get_wikimedia_images(topic: dict, action_type: str) -> list:
    """Return up to one relevant Wikimedia Commons image for the topic (empty list if none found)."""
    title = topic.get('title', '').strip().rstrip('.')
    if not title:
        return []

    significant = [w for w in re.findall(r"[A-Za-z][A-Za-z-]{3,}", title)]
    generic = {"introduction", "overview", "basics", "fundamentals", "concepts", "concept", "using", "with", "and", "the", "from", "into"}
    stems = tuple(w.lower()[:5] for w in significant if w.lower() not in generic)
    short = " ".join(significant[:3])
    queries = [f"{short or title} diagram", title] + ([short] if len(significant) > 3 else [])

    for query in queries:
        found = search_commons_image(query, stems)
        if found:
            return [{
                "id": f"img_{uuid.uuid4().hex[:10]}",
                "url": found["url"],
                "caption": f"{title} (Wikimedia Commons)",
                "alt": f"Illustration related to {title}",
                "source": "wikimedia",
                "sourcePage": found["page"],
                "isPlaceholder": False
            }]
    return []

@app.post("/generate-dynamic-slide")
def generate_dynamic_slide(
    session_id: str = Body(...),
    lecture_plan: dict = Body(...),
    student_goal: str = Body('gate'),
    knowledge_state: dict = Body({}),
    current_topic_index: int = Body(0),
    action_type: Optional[str] = Body('concept_explanation'),
    user_question: Optional[str] = Body(None),
    context: Optional[str] = Body('')
):
    """
    Generate dynamic slide with AI using three input streams:
    1. Student Goal & Exam Preference
    2. Knowledge State from Quiz Results
    3. Lecture Plan Content
    """
    action_type = action_type or 'concept_explanation'
    context = context or ''
    topics = lecture_plan.get('impTopics', [])

    if not topics:
        raise HTTPException(status_code=400, detail="Lecture plan has no topics")
    if current_topic_index < 0 or current_topic_index >= len(topics):
        raise HTTPException(status_code=400, detail=f"Invalid topic index: {current_topic_index}")

    current_topic = topics[current_topic_index]
    print(f"Slide request: action={action_type} goal={student_goal} topic={current_topic.get('title')}")

    if action_type == 'completion':
        context = context or ", ".join(t.get('title', '') for t in topics)

    collection = lecture_plan.get('sessionId')
    sources, earlier = [], []
    if action_type in ('concept_explanation', 'explain-more', 'ask-question', 'simplified-explanation'):
        query = user_question or f"{current_topic.get('title', '')} {current_topic.get('description', '')[:200]}"
        sources, earlier = retrieve_reference(collection, query, include_explanations=action_type in ('ask-question', 'explain-more'))

    reference = ""
    if sources:
        reference += "\n".join(f"[{i + 1}] {snippet(s['text'], 600)}" for i, s in enumerate(sources))
    if earlier:
        reference += "\n\nEARLIER EXPLANATIONS GIVEN TO THIS STUDENT (do not repeat them, build on them):\n"
        reference += "\n".join(f"- {snippet(e['text'], 300)}" for e in earlier)

    slide_content = generate_slide_with_ai(
        topic=current_topic,
        action_type=action_type,
        student_goal=student_goal,
        knowledge_state=knowledge_state or {},
        user_question=user_question,
        context=context,
        reference=reference
    )

    images = get_wikimedia_images(current_topic, action_type)

    slide = {
        "id": f"slide_{int(datetime.now().timestamp()*1000)}",
        "title": slide_content.get('title', current_topic.get('title', 'Learning Slide')),
        "type": action_type,
        "content": slide_content,
        "images": images,
        "metadata": {
            "generatedAt": datetime.now().isoformat(),
            "sessionId": session_id,
            "topicIndex": current_topic_index,
            "topicTitle": current_topic.get('title', ''),
            "studentGoal": student_goal,
            "knowledgeLevel": (knowledge_state or {}).get('knowledgeLevel', 'intermediate'),
            "actionType": action_type,
            "hasImages": len(images) > 0,
            "generatedBy": "AI",
            "aiModel": OPENROUTER_MODEL,
            "sources": [{"id": s["id"], "kind": s["kind"], "snippet": snippet(s["text"]), "score": s["score"]} for s in sources]
        }
    }

    # Remember what was explained so later answers can build on it
    if collection and action_type not in ('quiz-results', 'completion'):
        explanation = " ".join(filter(None, [
            slide_content.get('title', ''), slide_content.get('narrative', ''),
            " ".join(str(i) for i in slide_content.get('keyInsights', [])[:4])
        ]))
        rag_store.add(collection, [{
            "id": slide["id"], "text": explanation, "kind": "explanation",
            "meta": {"topicIndex": current_topic_index, "action": action_type}
        }])

    return {"success": True, "slide": slide, "message": f"AI-generated slide for {action_type}"}

@app.post("/analyze-lecture")
def analyze_lecture(
    pdf: Optional[UploadFile] = File(None),
    ppt: Optional[UploadFile] = File(None),
    audio: Optional[UploadFile] = File(None),
    yt_link: Optional[str] = None
):
    """Turn uploaded material into a lecture plan. Accepts a PDF, a PPTX, an audio file and/or a YouTube link (query parameter yt_link); a transcript on its own is enough. Saves the plan, indexes every chunk for retrieval and returns the plan."""
    if not pdf and not ppt and not audio and not yt_link:
        raise HTTPException(status_code=400, detail="Provide at least one of: pdf, ppt, audio or yt_link.")

    print("Extracting content...")
    pdf_text = extract_pdf_text(pdf) if pdf else ""
    ppt_text = extract_ppt_text(ppt) if ppt else ""
    transcript = None

    if audio:
        print("Transcribing audio...")
        transcript = transcribe_audio(audio)
    elif yt_link:
        print("Processing YouTube link...")
        transcript = process_yt_link(yt_link)

    # Slides/documents are the primary source; a transcript alone is enough when there are none.
    source_text = f"{pdf_text}\n{ppt_text}".strip() or (transcript or "")
    if not source_text:
        raise HTTPException(status_code=422, detail="No readable text could be extracted from the input.")

    print("Analyzing content...")
    analysis = analyze_text(source_text, transcript)

    print("Generating lecture plan...")
    plan = build_lecture_plan(analysis)

    save_lecture_plan(plan)
    try:
        index_lecture(plan, analysis, transcript, transcript_is_source=not (pdf_text or ppt_text))
    except Exception as e:  # retrieval is an enhancement; never fail the upload because of it
        print(f"Could not index lecture for retrieval: {e}")
    return plan

@app.post("/generate-diagnostic-mcq")
def generate_diagnostic_mcq(request: MCQRequest):
    """Generate diagnostic MCQs from lecture plan"""
    topics = request.lecture_plan.get('impTopics', [])
    if not topics:
        raise HTTPException(status_code=400, detail="No topics found in lecture plan")

    num_questions = max(1, min(request.num_questions, 30))
    print(f"Generating {num_questions} MCQs from {len(topics)} topics...")

    # Spread the questions across topics as evenly as possible
    plan = [topics[i % len(topics)] for i in range(num_questions)]
    plan.sort(key=lambda t: topics.index(t))

    def make(args):
        qid, topic = args
        content = topic.get('full_content') or topic.get('description', '')
        return create_diagnostic_question(qid, topic, content, topic.get('keyInsights', []), topics)

    with ThreadPoolExecutor(max_workers=4) as pool:
        questions = list(pool.map(make, enumerate(plan, start=1)))

    return {
        "success": True,
        "questions": questions,
        "metadata": {
            "generatedAt": datetime.now().isoformat(),
            "totalQuestions": len(questions),
            "topicsUsed": len(topics)
        }
    }

@app.post("/generate-topic-quiz")
def generate_topic_quiz(request: TopicQuizRequest):
    """Generate an adaptive quiz for one topic. 503 means the AI is unavailable (client falls back)."""
    quiz = generate_topic_quiz_with_llm(request)
    if not quiz:
        raise HTTPException(status_code=503, detail="AI quiz generation is unavailable")
    return {"success": True, "quiz": quiz}

class RagQuery(BaseModel):
    collection: str
    query: str
    k: int = 4

@app.post("/rag/query")
def rag_query(req: RagQuery):
    """Retrieve the most relevant lecture passages for a question."""
    results = rag_store.query(req.collection, req.query, k=max(1, min(req.k, 10)))
    return {"results": [{**r, "snippet": snippet(r["text"])} for r in results]}

@app.post("/analyze-quiz-results")
def analyze_quiz_results(request: AnswerAnalysisRequest):
    """Analyze quiz results and identify misconceptions"""
    questions = request.questions
    answers_by_id = {str(a.get('questionId')): a for a in request.answers}

    correct = 0
    misconceptions, weak_topics = [], []
    for question in questions:
        answer = answers_by_id.get(str(question.get('id')))
        if answer and answer.get('selectedAnswer') == question.get('correctAnswer'):
            correct += 1
            continue
        topic = (question.get('conceptsTested') or ['Unknown'])[0]
        weak_topics.append(topic)
        misconceptions.append({
            "topic": topic,
            "question": question.get('question'),
            "yourAnswer": answer.get('selectedAnswer') if answer else None,
            "correctAnswer": question.get('correctAnswer'),
            "explanation": question.get('explanation', ''),
            "type": "knowledge_gap"
        })

    score = (correct / len(questions)) * 100 if questions else 0
    recommendations = []
    if score < 60:
        recommendations.append("Review fundamental concepts from the lecture")
    if weak_topics:
        recommendations.append(f"Focus on: {', '.join(list(dict.fromkeys(weak_topics))[:3])}")
    if score >= 80:
        recommendations.append("Strong understanding! Move to advanced topics")

    return {
        "success": True,
        "score": round(score, 1),
        "correct": correct,
        "total": len(questions),
        "misconceptions": misconceptions,
        "weakTopics": list(dict.fromkeys(weak_topics)),
        "recommendations": recommendations,
        "knowledgeState": {
            "overallMastery": score,
            "conceptualUnderstanding": score,
            "needsReview": score < 70
        }
    }

@app.post("/save-quiz-results")
def save_quiz_results_endpoint(request: SaveResultsRequest):
    """Save quiz results to disk"""
    try:
        save_quiz_results({
            "session_id": request.session_id,
            "lecture_id": request.lecture_id,
            "score": request.score,
            "correct": request.correct,
            "total": request.total,
            "questions": request.questions,
            "answers": request.answers,
            "misconceptions": request.misconceptions,
            "weak_topics": request.weak_topics,
            "recommendations": request.recommendations,
            "timestamp": datetime.now().isoformat()
        })
    except OSError as e:
        raise HTTPException(status_code=500, detail=f"Failed to save results: {e}")

    return {"success": True, "message": "Quiz results saved successfully", "session_id": request.session_id}

# =========================================================
# HEALTH CHECK / LOOKUP
# =========================================================
@app.get("/")
def root():
    """Service banner: version, whether an AI key is configured, and the list of endpoints."""
    return {
        "message": "Study Plan Generator API",
        "version": "2.1",
        "ai_configured": bool(OPENROUTER_API_KEY),
        "endpoints": [
            "/analyze-lecture",
            "/generate-diagnostic-mcq",
            "/generate-dynamic-slide",
            "/generate-topic-quiz",
            "/analyze-quiz-results",
            "/save-quiz-results",
            "/get-lecture-plan",
            "/rag/query",
            "/classroom"
        ]
    }

@app.get("/health")
def health():
    """Liveness probe."""
    return {"status": "healthy"}

@app.get("/get-lecture-plan")
def get_lecture_plan(
    sessionId: Optional[str] = None,
    lectureId: Optional[str] = None,
    latest: Optional[bool] = False
):
    """Return the lecture plan for a session/lecture id, or the most recent plan when `latest` is set."""
    plans = _read_json_list(LECTURE_PLANS_FILE)
    if not plans:
        raise HTTPException(status_code=404, detail="No lecture plans found")

    for plan in plans:
        if (sessionId and plan.get('sessionId') == sessionId) or (lectureId and plan.get('lectureId') == lectureId):
            return plan

    if latest or not (sessionId or lectureId):
        return plans[0]

    raise HTTPException(status_code=404, detail="Lecture plan not found for that id")

# =========================================================
# SERVER STARTUP
# =========================================================
if __name__ == "__main__":
    import uvicorn
    print(f"Study Plan Generator API: http://localhost:{PORT}  (docs: http://localhost:{PORT}/docs)")
    uvicorn.run(app, host="127.0.0.1", port=PORT)
