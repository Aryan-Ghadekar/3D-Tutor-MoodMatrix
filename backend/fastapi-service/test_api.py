"""API tests for the FastAPI service. External services (OpenRouter, Wikimedia, YouTube, ML models)
are mocked, so these run offline:  pip install pytest httpx  &&  pytest backend/fastapi-service
"""
import io
import json
import os
import tempfile

os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ.pop("OPENROUTER_API_KEY", None)

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

import main

client = TestClient(main.app)

LECTURE_TEXT = [
    "Binary search trees keep keys ordered so that lookup, insertion and deletion take logarithmic time on balanced trees. "
    "Important: an unbalanced tree degrades to a linked list. Remember that rotations restore balance in AVL trees. " * 6,
    "Hash tables map keys to buckets using a hash function. A good hash function spreads keys uniformly. "
    "Collisions are resolved by chaining or open addressing. Note that the load factor controls resizing. " * 6,
    "Graph traversal uses breadth first search with a queue or depth first search with a stack. "
    "The key idea of Dijkstra's algorithm is greedy selection of the closest unvisited vertex. " * 6,
]


@pytest.fixture(autouse=True)
def no_external_calls(monkeypatch):
    def fail(*args, **kwargs):
        raise Exception("LLM disabled in tests")
    monkeypatch.setattr(main, "llm_message", fail)
    monkeypatch.setattr(main, "search_commons_image", lambda q, s=(): None)
    monkeypatch.setattr(main, "get_semantic_model", lambda: None)


def make_pptx_bytes():
    prs = Presentation()
    for text in LECTURE_TEXT:
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        box = slide.shapes.add_textbox(0, 0, 100, 100)
        box.text_frame.text = text
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_health_and_root():
    assert client.get("/health").json() == {"status": "healthy"}
    assert "/analyze-lecture" in client.get("/").json()["endpoints"]


@pytest.mark.parametrize("url,expected", [
    ("https://www.youtube.com/watch?v=dQw4w9WgXcQ&t=3s", "dQw4w9WgXcQ"),
    ("https://youtu.be/dQw4w9WgXcQ?si=abc", "dQw4w9WgXcQ"),
    ("https://www.youtube.com/shorts/dQw4w9WgXcQ", "dQw4w9WgXcQ"),
    ("https://example.com/video", None),
])
def test_extract_video_id(url, expected):
    assert main.extract_video_id(url) == expected


def test_analyze_lecture_requires_input():
    assert client.post("/analyze-lecture").status_code == 400


def test_analyze_lecture_from_pptx_without_ai():
    files = {"ppt": ("lecture.pptx", make_pptx_bytes(),
                     "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    r = client.post("/analyze-lecture", files=files)
    assert r.status_code == 200, r.text
    plan = r.json()
    assert plan["impTopics"], "plan should have topics"
    for topic in plan["impTopics"]:
        assert topic["title"] and "failed" not in topic["title"].lower()
        assert all("failed" not in s.lower() for s in topic["keyInsights"])
    assert "saved_path" not in plan
    # the plan was persisted and can be fetched back
    got = client.get("/get-lecture-plan", params={"sessionId": plan["sessionId"]})
    assert got.status_code == 200 and got.json()["lectureId"] == plan["lectureId"]


def test_analyze_lecture_from_youtube_transcript_only(monkeypatch):
    monkeypatch.setattr(main, "process_yt_link", lambda link: " ".join(LECTURE_TEXT))
    r = client.post("/analyze-lecture", params={"yt_link": "https://youtu.be/dQw4w9WgXcQ"})
    assert r.status_code == 200, r.text
    assert r.json()["impTopics"]


def test_analyze_lecture_topics_keep_lecture_order(monkeypatch):
    monkeypatch.setattr(main, "process_yt_link", lambda link: " ".join(LECTURE_TEXT) * 3)
    plan = client.post("/analyze-lecture", params={"yt_link": "https://youtu.be/dQw4w9WgXcQ"}).json()
    ids = [t["id"] for t in plan["impTopics"]]
    assert ids == [f"topic-{i + 1}" for i in range(len(ids))]


def test_get_lecture_plan_lookup_rules(monkeypatch):
    monkeypatch.setattr(main, "process_yt_link", lambda link: " ".join(LECTURE_TEXT))
    plan = client.post("/analyze-lecture", params={"yt_link": "https://youtu.be/dQw4w9WgXcQ"}).json()
    assert client.get("/get-lecture-plan", params={"sessionId": plan["sessionId"]}).json()["lectureId"] == plan["lectureId"]
    # unknown id is a 404 unless the caller explicitly asks for the latest plan
    assert client.get("/get-lecture-plan", params={"sessionId": "nope"}).status_code == 404
    latest = client.get("/get-lecture-plan", params={"sessionId": "nope", "latest": True})
    assert latest.status_code == 200 and latest.json()["lectureId"] == plan["lectureId"]


def sample_plan():
    return {"impTopics": [
        {"id": "topic-1", "title": "Binary Search Trees", "description": "Ordered trees give logarithmic lookup. They need balancing.",
         "keyInsights": ["Balanced trees give O(log n) lookups"], "importance": "high"},
        {"id": "topic-2", "title": "Hash Tables", "description": "Hash tables map keys to buckets. Collisions need resolving.",
         "keyInsights": ["Chaining resolves collisions"], "importance": "medium"},
        {"id": "topic-3", "title": "Graph Traversal", "description": "BFS uses a queue. DFS uses a stack.",
         "keyInsights": ["BFS finds shortest paths in unweighted graphs"], "importance": "medium"},
    ]}


def test_diagnostic_mcq_template_fallback_is_valid():
    r = client.post("/generate-diagnostic-mcq", json={"lecture_plan": sample_plan(), "num_questions": 6})
    assert r.status_code == 200, r.text
    questions = r.json()["questions"]
    assert len(questions) == 6
    assert len({q["id"] for q in questions}) == 6
    for q in questions:
        assert [o["id"] for o in q["options"]] == ["A", "B", "C", "D"]
        correct = [o for o in q["options"] if o["isCorrect"]]
        assert len(correct) == 1 and correct[0]["id"] == q["correctAnswer"]
        assert len({o["text"] for o in q["options"]}) == 4, "options must be distinct"


def test_normalize_mcq_rejects_bad_and_shuffles():
    assert main.normalize_mcq({"question": "q", "options": []}) is None
    bad = {"question": "q", "options": [{"text": str(i), "isCorrect": True} for i in range(4)]}
    assert main.normalize_mcq(bad) is None
    good = {"question": "q", "options": [{"text": "right", "isCorrect": True}] +
            [{"text": f"w{i}", "isCorrect": False} for i in range(3)]}
    positions = {next(o["id"] for o in main.normalize_mcq(good)["options"] if o["isCorrect"]) for _ in range(60)}
    assert len(positions) > 1, "the correct answer must not always land in the same slot"


def quiz_request(**kw):
    body = {"topic": sample_plan()["impTopics"][0], "difficulty": "easy", "studentGoal": "gate",
            "numberOfQuestions": 2, "sessionId": "s1"}
    body.update(kw)
    return body


def test_topic_quiz_unavailable_without_ai():
    assert client.post("/generate-topic-quiz", json=quiz_request()).status_code == 503


def test_topic_quiz_with_ai(monkeypatch):
    def q(text):
        return {"question": text, "type": "conceptual", "explanation": "because",
                "options": [{"text": "right", "isCorrect": True}] + [{"text": f"w{i}", "isCorrect": False} for i in range(3)]}
    monkeypatch.setattr(main, "llm_message", lambda *a, **k: "```json\n" + json.dumps({"questions": [q("Q1?"), q("Q2?")]}) + "\n```")
    r = client.post("/generate-topic-quiz", json=quiz_request())
    assert r.status_code == 200, r.text
    quiz = r.json()["quiz"]
    assert len(quiz["questions"]) == 2
    for question in quiz["questions"]:
        assert question["correctAnswer"] in "ABCD" and question["text"] == question["question"]


SLIDE = {"title": "T", "narrative": "n", "keyInsights": ["a"], "analogies": [], "realWorldScenario": None,
         "interactiveElement": "?", "examRelevance": "x", "commonTraps": [], "memoryHook": "m"}


def slide_request(**kw):
    body = {"session_id": "s1", "lecture_plan": sample_plan(), "student_goal": "gate", "knowledge_state": {},
            "current_topic_index": 1}
    body.update(kw)
    return body


@pytest.mark.parametrize("action", ["concept_explanation", "explain-more", "ask-question", "quiz-results", "completion", None])
def test_dynamic_slide_all_actions(monkeypatch, action):
    monkeypatch.setattr(main, "llm_message", lambda *a, **k: json.dumps(SLIDE))
    body = slide_request(action_type=action, user_question="why?" if action == "ask-question" else None)
    r = client.post("/generate-dynamic-slide", json=body)
    assert r.status_code == 200, r.text
    slide = r.json()["slide"]
    assert slide["type"] == (action or "concept_explanation")
    assert slide["metadata"]["topicIndex"] == 1


def test_dynamic_slide_bad_index_and_ai_failure():
    assert client.post("/generate-dynamic-slide", json=slide_request(current_topic_index=9)).status_code == 400
    assert client.post("/generate-dynamic-slide", json=slide_request()).status_code == 502


def test_dynamic_slide_accepts_weak_topics_and_misconceptions(monkeypatch):
    monkeypatch.setattr(main, "llm_message", lambda *a, **k: json.dumps(SLIDE))
    ks = {"knowledgeLevel": "beginner", "averageScore": 40, "weakTopics": ["Hash Tables"],
          "misconceptions": [{"topic": "Hash Tables", "misconception": "x"}, "junk"]}
    assert client.post("/generate-dynamic-slide", json=slide_request(knowledge_state=ks, action_type="quiz-results")).status_code == 200


def test_parse_llm_json_variants():
    assert main.parse_llm_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert main.parse_llm_json('Here you go: {"a": 1} thanks') == {"a": 1}
    assert main.parse_llm_json('["x", "y"]') == ["x", "y"]
    with pytest.raises(ValueError):
        main.parse_llm_json("no json here")


def test_save_quiz_results_round_trip():
    body = {"session_id": "s", "lecture_id": "l", "score": 50.0, "correct": 1, "total": 2,
            "questions": [], "answers": [], "misconceptions": [], "weak_topics": [], "recommendations": []}
    assert client.post("/save-quiz-results", json=body).json()["success"] is True
    saved = json.loads(main.QUIZ_RESULTS_FILE.read_text(encoding="utf-8"))
    assert saved[0]["session_id"] == "s"


def test_analyze_quiz_results_matches_answers_by_id():
    questions = [{"id": "q1", "correctAnswer": "A", "conceptsTested": ["T1"]},
                 {"id": "q2", "correctAnswer": "B", "conceptsTested": ["T2"]}]
    answers = [{"questionId": "q2", "selectedAnswer": "B"}]  # q1 unanswered
    r = client.post("/analyze-quiz-results", json={"session_id": "s", "questions": questions, "answers": answers}).json()
    assert r["correct"] == 1 and r["weakTopics"] == ["T1"]
